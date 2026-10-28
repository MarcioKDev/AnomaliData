"""Leitura e validação de arquivos tabulares de compras."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable, Mapping

from anomalidata.dominio import DadoInvalidoError, RegistroCompra


COLUNAS_OBRIGATORIAS = (
    "id_registro",
    "codigo_produto",
    "fornecedor",
    "quantidade",
    "preco_unitario",
    "data_compra",
)


@dataclass(frozen=True, slots=True)
class ResultadoImportacao:
    registros: tuple[RegistroCompra, ...]
    total_lidos: int
    linhas_duplicadas: int


def _decimal(valor: object, campo: str, linha: int) -> Decimal:
    texto = str(valor).strip().replace(",", ".")
    try:
        return Decimal(texto)
    except (InvalidOperation, ValueError) as error:
        raise DadoInvalidoError(
            f"Linha {linha}: {campo} deve conter um número válido"
        ) from error


def _data(valor: object, linha: int) -> date:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()
    try:
        return date.fromisoformat(texto)
    except ValueError as error:
        raise DadoInvalidoError(
            f"Linha {linha}: data_compra deve usar o formato AAAA-MM-DD"
        ) from error


def _validar_colunas(colunas: Iterable[object]) -> tuple[str, ...]:
    nomes = tuple(str(coluna).strip() for coluna in colunas if coluna is not None)
    ausentes = [coluna for coluna in COLUNAS_OBRIGATORIAS if coluna not in nomes]
    if ausentes:
        raise DadoInvalidoError(
            "Colunas obrigatórias ausentes: " + ", ".join(ausentes)
        )
    return nomes


def _converter_linhas(
    linhas: Iterable[Mapping[str, object]],
) -> ResultadoImportacao:
    registros: list[RegistroCompra] = []
    assinaturas: set[tuple[str, ...]] = set()
    ids_vistos: set[str] = set()
    duplicadas = 0
    total = 0

    for numero, linha in enumerate(linhas, start=2):
        if not any(valor not in (None, "") for valor in linha.values()):
            continue
        total += 1
        assinatura = tuple(
            str(linha.get(coluna, "")).strip() for coluna in COLUNAS_OBRIGATORIAS
        )
        if assinatura in assinaturas:
            duplicadas += 1
            continue
        assinaturas.add(assinatura)
        id_registro = str(linha.get("id_registro", "")).strip()
        if id_registro in ids_vistos:
            raise DadoInvalidoError(
                f"Linha {numero}: id_registro duplicado com conteúdo diferente"
            )
        ids_vistos.add(id_registro)
        try:
            registro = RegistroCompra(
                id_registro=id_registro,
                codigo_produto=str(linha["codigo_produto"]),
                fornecedor=str(linha["fornecedor"]),
                quantidade=_decimal(linha["quantidade"], "quantidade", numero),
                preco_unitario=_decimal(
                    linha["preco_unitario"],
                    "preco_unitario",
                    numero,
                ),
                data_compra=_data(linha["data_compra"], numero),
            )
        except DadoInvalidoError as error:
            if str(error).startswith("Linha "):
                raise
            raise DadoInvalidoError(f"Linha {numero}: {error}") from error
        registros.append(registro)

    if not registros:
        raise DadoInvalidoError("O arquivo não contém registros válidos")
    return ResultadoImportacao(tuple(registros), total, duplicadas)


def _carregar_csv(caminho: Path) -> ResultadoImportacao:
    with caminho.open(encoding="utf-8-sig", newline="") as arquivo:
        amostra = arquivo.read(4096)
        arquivo.seek(0)
        try:
            dialeto = csv.Sniffer().sniff(amostra, delimiters=",;\t")
        except csv.Error:
            dialeto = csv.excel
        leitor = csv.DictReader(arquivo, dialect=dialeto)
        if leitor.fieldnames is None:
            raise DadoInvalidoError("O arquivo CSV não possui cabeçalho")
        _validar_colunas(leitor.fieldnames)
        return _converter_linhas(leitor)


def _carregar_xlsx(caminho: Path) -> ResultadoImportacao:
    from openpyxl import load_workbook

    workbook = load_workbook(caminho, read_only=True, data_only=True)
    try:
        planilha = workbook.active
        linhas = planilha.iter_rows(values_only=True)
        try:
            cabecalho = next(linhas)
        except StopIteration as error:
            raise DadoInvalidoError("A planilha não possui cabeçalho") from error
        colunas = _validar_colunas(cabecalho)
        mapas = (dict(zip(colunas, valores, strict=False)) for valores in linhas)
        return _converter_linhas(mapas)
    finally:
        workbook.close()


def carregar_arquivo(caminho: str | Path) -> ResultadoImportacao:
    """Carrega CSV ou XLSX e devolve somente linhas únicas e válidas."""

    arquivo = Path(caminho)
    if not arquivo.is_file():
        raise DadoInvalidoError("Arquivo de entrada não encontrado")
    extensao = arquivo.suffix.lower()
    if extensao == ".csv":
        return _carregar_csv(arquivo)
    if extensao == ".xlsx":
        return _carregar_xlsx(arquivo)
    raise DadoInvalidoError("Formato não suportado; use CSV ou XLSX")
